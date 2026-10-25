"""Adapters that read the patch list from its configured sources."""

from __future__ import annotations

import contextlib
import csv
import email
import email.utils
import re
from dataclasses import dataclass, field
from email.header import decode_header, make_header
from pathlib import Path

from patchmig.config import ProjectConfig
from patchmig.config.schema import (
    DiffFolderSource,
    GitRangeSource,
    JiraQuerySource,
    SpreadsheetSource,
)
from patchmig.gitutil import (
    GitError,
    commit_info,
    file_at,
    is_ancestor,
    range_log,
    resolve,
    rev_list,
)
from patchmig.ingest.codehosts import (
    CodeHostError,
    CodeRef,
    HostClient,
    find_refs,
    host_for,
    is_known,
)
from patchmig.ingest.diffparse import parse_diff
from patchmig.ingest.manifest import project_for_path
from patchmig.ingest.normalize import RawPatch
from patchmig.jira_client import JiraError, JiraReader, Ticket
from patchmig.knowledge.history import has_range, is_shallow

SPREADSHEET_FIELDS = ("commit", "project", "jira", "subject")


@dataclass
class SourceContext:
    config: ProjectConfig
    projects: list[str]  # project paths in the source tree
    jira: JiraReader | None = None
    has_tree: bool = True  # False: jira_query without a source checkout
    warnings: list[str] = field(default_factory=list)
    # Every ticket a JQL source found, with or without a patch: each one is kept (ingest).
    tickets: dict[str, Ticket] = field(default_factory=dict)
    # What failed while reading a ticket (Jira or a code host refused): by ticket key.
    problems: dict[str, list[str]] = field(default_factory=dict)
    _range: dict[str, list[tuple[str, str]]] = field(default_factory=dict, repr=False)
    _remote: dict[str, list[dict]] = field(default_factory=dict, repr=False)

    @property
    def tree(self) -> Path:
        return self.config.trees.source.path

    @property
    def base(self) -> str:
        return self.config.trees.source.base_ref

    @property
    def head(self) -> str:
        return self.config.trees.source.head_ref  # validated non-None

    @property
    def downloads(self) -> Path:
        """Where diffs fetched from Jira or a code host are kept (patches point at them)."""
        return self.config.workdir / "downloads"

    def warn(self, message: str) -> None:
        self.warnings.append(message)

    def problem(self, keys: str | list[str], message: str) -> None:
        """A warning that means the ticket was not read completely: import it again later."""
        self.warn(message)
        for key in [keys] if isinstance(keys, str) else keys:
            self.problems.setdefault(key, []).append(message)

    def remote_links(self, key: str) -> list[dict]:
        """A ticket's remote links (web links), read once; empty when Jira refuses."""
        if key not in self._remote:
            try:
                self._remote[key] = self.jira.remote_links(key)
            except JiraError as e:
                self.problem(key, f"jira_query: {key} remote links: {e}")
                self._remote[key] = []
        return self._remote[key]

    def range_log(self, project: str) -> list[tuple[str, str]]:
        """(sha, message) of the project's OEM-range commits, oldest first; read once.

        Empty for a project without the range (no such repo, or base/head missing in it).
        """
        if project not in self._range:
            repo = self.tree / project
            self._range[project] = []
            if repo.is_dir():
                with contextlib.suppress(GitError):
                    self._range[project] = range_log(repo, self.base, self.head)
        return self._range[project]


def from_git_range(src: GitRangeSource, ctx: SourceContext) -> list[RawPatch]:
    projects = ctx.projects if src.projects == "all" else src.projects
    patches: list[RawPatch] = []
    for project in projects:
        repo = ctx.tree / project
        if project not in ctx.projects:
            ctx.warn(f"git_range: project {project} is not in the source tree")
            continue
        if is_shallow(repo) and not has_range(repo, ctx.base, ctx.head):
            # base..head in a shallow clone silently lists only the commits it holds
            ctx.warn(f"git_range: {project}: shallow checkout without the OEM range; "
                     "run migrate history first")  # fmt: skip
            continue
        try:
            shas = rev_list(repo, ctx.base, ctx.head)
        except GitError as e:
            ctx.warn(f"git_range: {project}: {e.stderr}")
            continue
        patches += [
            RawPatch.from_commit(project, commit_info(repo, sha), "git_range") for sha in shas
        ]
    return patches


def _find_commit(ctx: SourceContext, sha: str, project: str | None) -> tuple[str, str] | None:
    candidates = [project] if project else ctx.projects
    for p in candidates:
        repo = ctx.tree / p
        if repo.exists() and (full := resolve(repo, sha)):
            return p, full
    return None


def _decode(value: str | None) -> str:
    return str(make_header(decode_header(value))) if value else ""


def _parse_patch_file(path: Path) -> tuple[str, str, str | None, str]:
    """Return (subject, message, author, diff) of a format-patch mail or a plain diff."""
    text = path.read_text(errors="replace")
    if text.startswith("From ") or text.startswith("From:") or "\nSubject:" in text[:2000]:
        msg = email.message_from_string(text)
        subject = re.sub(r"^\[PATCH[^\]]*\]\s*", "", _decode(msg["Subject"]).replace("\n", " "))
        body = msg.get_payload(decode=False)
        body = body if isinstance(body, str) else ""
        description, _, diff = body.partition("\n---\n")
        if "diff --git" not in diff:  # no '---' separator: split at the first diff header
            description, _, rest = body.partition("diff --git")
            diff = "diff --git" + rest if rest else ""
        else:
            diff = diff[diff.find("diff --git") :] if "diff --git" in diff else ""
        diff = re.sub(r"\n-- \n[^\n]*\n*$", "\n", diff)  # format-patch signature
        message = f"{subject}\n\n{description.strip()}".strip()
        return subject, message, _decode(msg["From"]) or None, diff
    subject = path.stem
    return subject, subject, None, text


_PATH_LINE = re.compile(
    r"^(diff --git a/|--- a/|\+\+\+ b/|rename from |rename to |copy from |copy to )"
)


def strip_project_prefix(diff: str, project: str) -> str:
    """Rewrite tree-relative paths in a diff's header lines to project-relative ones."""
    prefix = project.rstrip("/") + "/"
    out = []
    for line in diff.splitlines(keepends=True):
        if m := _PATH_LINE.match(line):
            rest = line[m.end() :]
            if line.startswith("diff --git"):
                rest = rest.replace(f" b/{prefix}", " b/", 1)
            line = m.group(1) + rest.removeprefix(prefix)
        out.append(line)
    return "".join(out)


def from_diff_folder(src: DiffFolderSource, ctx: SourceContext) -> list[RawPatch]:
    if not src.path.is_dir():
        ctx.warn(f"diff_folder: {src.path} is not a folder")
        return []
    patches: list[RawPatch] = []
    for path in sorted(src.path.rglob("*")):
        if path.suffix not in (".patch", ".diff") or not path.is_file():
            continue
        subject, message, author, diff = _parse_patch_file(path)
        if not diff.strip():
            ctx.warn(f"diff_folder: {path.name} contains no diff")
            continue
        folder = path.parent.relative_to(src.path).as_posix()
        project = folder if folder in ctx.projects else None
        if project is None:
            # Tree-relative paths inside the diff: map by longest project prefix.
            files = parse_diff(diff).files
            owners = {project_for_path(f.path, ctx.projects) for f in files}
            if len(owners) == 1 and None not in owners:
                project = owners.pop()
                diff = strip_project_prefix(diff, project)
        if project is None:
            ctx.warn(
                f"diff_folder: cannot tell which project {path.name} belongs to; put it in a "
                "subfolder named after the project path"
            )
            continue
        patches.append(
            RawPatch(
                project=project,
                subject=subject,
                message=message,
                diff=diff,
                source="diff_folder",
                diff_path=str(path),
                author=author,
            )
        )
    return patches


def _rows(src: SpreadsheetSource) -> list[dict[str, str]]:
    if src.path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(src.path, read_only=True, data_only=True)
        ws = wb[src.sheet] if src.sheet else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
        return [
            {h: "" if v is None else str(v).strip() for h, v in zip(header, row, strict=False)}
            for row in rows
            if any(v is not None for v in row)
        ]
    with src.path.open(newline="", encoding="utf-8-sig") as f:
        return [
            {k.strip(): (v or "").strip() for k, v in row.items() if k} for row in csv.DictReader(f)
        ]


def mentions(key: str) -> re.Pattern:
    """A Jira key as a whole word: AUTO-12 must not match AUTO-123 or XAUTO-12."""
    return re.compile(rf"(?<![\w-]){re.escape(key)}(?!\d)")


def _jira_commits(ctx: SourceContext, key: str, source: str) -> list[RawPatch]:
    """OEM-range commits whose message mentions the key (each project's log is read once)."""
    pattern = mentions(key)
    found: list[RawPatch] = []
    for project in ctx.projects:
        repo = ctx.tree / project
        for sha, message in ctx.range_log(project):
            if not pattern.search(message):
                continue
            raw = RawPatch.from_commit(project, commit_info(repo, sha), source)
            raw.jira_keys.append(key)
            found.append(raw)
    return found


def from_spreadsheet(src: SpreadsheetSource, ctx: SourceContext) -> list[RawPatch]:
    try:
        rows = _rows(src)
    except (OSError, KeyError, ValueError) as e:
        ctx.warn(f"spreadsheet: cannot read {src.path}: {e}")
        return []
    col = {name: src.columns.get(name, name) for name in SPREADSHEET_FIELDS}
    patches: list[RawPatch] = []
    for n, row in enumerate(rows, start=2):
        sha = row.get(col["commit"], "")
        project = row.get(col["project"], "") or None
        keys = re.findall(r"[A-Z][A-Z0-9]+-\d+", row.get(col["jira"], ""))
        if sha:
            hit = _find_commit(ctx, sha, project)
            if hit is None:
                ctx.warn(f"spreadsheet row {n}: commit {sha} not found in the source tree")
                continue
            raw = RawPatch.from_commit(
                hit[0], commit_info(ctx.tree / hit[0], hit[1]), "spreadsheet"
            )
            raw.jira_keys += keys
            patches.append(raw)
        elif keys:
            found = [p for key in keys for p in _jira_commits(ctx, key, "spreadsheet")]
            if not found:
                ctx.warn(f"spreadsheet row {n}: no commit mentions {', '.join(keys)}")
            patches += found
        else:
            ctx.warn(f"spreadsheet row {n}: neither a commit nor a Jira key")
    return patches


def from_jira_query(src: JiraQuerySource, ctx: SourceContext) -> list[RawPatch]:
    """Patches of the tickets a JQL query finds, from every place a ticket can carry them:

    - commits in the source range that mention the ticket key;
    - .diff/.patch attachments of the ticket;
    - commit and PR links (GitHub, GitLab, Bitbucket) in the ticket and in the tickets it
      links to (the implementation tickets), in their text, comments and remote links.
    """
    if ctx.jira is None:
        ctx.warn("jira_query: no Jira connection available; skipped")
        return []
    patches: list[RawPatch] = []
    for ticket in ctx.jira.search(src.jql):
        patches += _ticket_patches(ctx, ticket, src)
    return patches


def _ticket_patches(ctx: SourceContext, ticket: Ticket, src: JiraQuerySource) -> list[RawPatch]:
    ctx.tickets[ticket.key] = ticket
    found = _jira_commits(ctx, ticket.key, "jira_query")
    if src.attachments:
        found += _attachment_patches(ctx, ticket, src.max_attachment_kb * 1024)
    if src.linked_tickets:
        found += _linked_patches(ctx, ticket)
    if not found:
        ctx.warn(f"jira_query: no patch found for {ticket.key} (commits, attachments, "
                 "links); the ticket is kept without one")  # fmt: skip
    return found


def from_ticket(ctx: SourceContext, key: str) -> list[RawPatch]:
    """One ticket, read again (or for the first time) as a JQL source reads its tickets.

    The options are those of the first jira_query source, else the defaults. Without a Jira
    connection only the commits that mention the key are found.
    """
    if ctx.jira is None:
        return _jira_commits(ctx, key, "jira_query")
    try:
        ticket = ctx.jira.get(key)
    except JiraError as e:
        ctx.problem(key, f"jira: {key}: {e}")
        return _jira_commits(ctx, key, "jira_query")
    if ticket is None:
        ctx.warn(f"jira: {key} not found")
        return _jira_commits(ctx, key, "jira_query")
    src = next((s for s in ctx.config.patch_sources if isinstance(s, JiraQuerySource)), None)
    return _ticket_patches(ctx, ticket, src or JiraQuerySource(type="jira_query", jql=""))


def from_additions(ctx: SourceContext, key: str, added: list[dict]) -> list[RawPatch]:
    """Patches in what an engineer added to a ticket: commit and PR links, .diff/.patch files.

    Other links and files are context for the analysis only.
    """
    found: list[RawPatch] = []
    for a in added:
        if a.get("removed_at"):
            continue
        if a["kind"] == "link":
            for ref in find_refs(a["url"], ctx.config.hosts()):
                found += _ref_patches(ctx, ref, [key])
        elif a["kind"] == "file" and a["filename"].endswith((".diff", ".patch")):
            path = Path(a["path"])
            if not path.is_file():
                ctx.warn(f"added to {key}: {a['filename']} is gone from {path.parent}")
                continue
            subject, message, author, diff = _parse_patch_file(path)
            if subject == path.stem and (ticket := ctx.tickets.get(key)):
                subject, message = ticket.summary, f"{ticket.summary}\n\n{ticket.description}"
            raw = _diff_patch(ctx, diff, f"{key}/added/{path.name}", subject, message.strip(),
                              "added", [key])  # fmt: skip
            if raw:
                raw.author = author
                found.append(raw)
    for raw in found:
        raw.source = "added"
    return found


def _in_range(ctx: SourceContext, project: str, sha: str) -> bool:
    repo = ctx.tree / project
    return is_ancestor(repo, sha, ctx.head) and not is_ancestor(repo, sha, ctx.base)


def _project_for_diff(ctx: SourceContext, diff: str, hint: str | None) -> tuple[str, str] | None:
    """(project, project-relative diff): from a repo mapping, tree paths, or existing files."""
    if hint in ctx.projects or (hint and not ctx.has_tree):
        return hint, diff
    files = parse_diff(diff).files
    if not files:
        return None
    owners = {project_for_path(f.path, ctx.projects) for f in files}
    if len(owners) == 1 and None not in owners:
        project = owners.pop()
        return project, strip_project_prefix(diff, project)
    if not ctx.has_tree:
        return None
    existing = [f.old_path or f.path for f in files if f.status != "A"] or [f.path for f in files]
    matches = [
        p for p in ctx.projects
        if all(file_at(ctx.tree / p, ctx.head, path) is not None for path in existing[:5])
    ]  # fmt: skip
    return (matches[0], diff) if len(matches) == 1 else None


def _diff_patch(
    ctx: SourceContext,
    diff: str,
    name: str,
    subject: str,
    message: str,
    source: str,
    keys: list[str],
    hint: str | None = None,
) -> RawPatch | None:
    if "diff --git" not in diff and "\n+++ " not in diff:
        return None
    where = _project_for_diff(ctx, diff, hint)
    if where is None:
        ctx.warn(f"{source}: cannot tell which project {name} belongs to (set a repo_map)")
        return None
    project, diff = where
    path = ctx.downloads / re.sub(r"[^\w.-]+", "_", name)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(diff)
    return RawPatch(project=project, subject=subject, message=message, diff=diff, source=source,
                    diff_path=str(path), jira_keys=list(keys))  # fmt: skip


def _attachment_patches(ctx: SourceContext, ticket: Ticket, max_bytes: int) -> list[RawPatch]:
    found = []
    for a in ticket.attachments:
        if not a["filename"].endswith((".diff", ".patch")):
            continue
        try:
            data = ctx.jira.download(a["url"], max_bytes).decode(errors="replace")
        except JiraError as e:
            ctx.problem(ticket.key, f"jira_query: {ticket.key} attachment {a['filename']}: {e}")
            continue
        tmp = attachment_path(ctx.downloads, ticket.key, a["filename"])
        tmp.parent.mkdir(parents=True, exist_ok=True)
        tmp.write_text(data)
        subject, message, author, diff = _parse_patch_file(tmp)
        if subject == tmp.stem:  # a plain diff: the ticket says what it is
            subject, message = ticket.summary, f"{ticket.summary}\n\n{ticket.description}".strip()
        raw = _diff_patch(ctx, diff, f"{ticket.key}/{a['filename']}", subject, message,
                          "jira_attachment", [ticket.key])  # fmt: skip
        if raw:
            raw.author = author
            found.append(raw)
    return found


def attachment_path(downloads: Path, key: str, filename: str) -> Path:
    """Where a ticket's attachment is kept: its name made safe, inside the ticket's folder."""
    name = re.sub(r"[^\w.-]+", "_", Path(filename).name).lstrip(".") or "attachment"
    return downloads / re.sub(r"[^\w.-]+", "_", key) / name


def _ticket_text(ctx: SourceContext, ticket: Ticket) -> str:
    remote = ctx.remote_links(ticket.key)
    return "\n".join([ticket.description, *ticket.comments, *(r["url"] for r in remote)])


def _linked_patches(ctx: SourceContext, ticket: Ticket) -> list[RawPatch]:
    found: list[RawPatch] = []
    sources = [(ticket, [ticket.key])]
    for key in ticket.issue_links:
        try:
            linked = ctx.jira.get(key)
        except JiraError as e:
            ctx.problem(ticket.key, f"jira_query: linked ticket {key}: {e}")
            continue
        if linked:
            sources.append((linked, [ticket.key, key]))
    for t, keys in sources:
        for ref in find_refs(_ticket_text(ctx, t), ctx.config.hosts()):
            found += _ref_patches(ctx, ref, keys)
    return found


def _local_commit(ctx: SourceContext, sha: str, project: str | None) -> RawPatch | None:
    if not ctx.has_tree:
        return None
    hit = _find_commit(ctx, sha, project)
    if hit is None or not _in_range(ctx, *hit):
        return None
    return RawPatch.from_commit(hit[0], commit_info(ctx.tree / hit[0], hit[1]), "jira_link")


def _ref_patches(ctx: SourceContext, ref: CodeRef, keys: list[str]) -> list[RawPatch]:
    hosts = ctx.config.hosts()
    host = host_for(ref, hosts)
    hint = getattr(host, "repo_map", {}).get(ref.repo)
    found: list[RawPatch] = []
    try:
        if ref.sha and (raw := _local_commit(ctx, ref.sha, hint)):
            found.append(raw)  # a commit link needs no host: the commit is in the source tree
        elif not is_known(ref, hosts):
            ctx.warn(f"jira_query: {ref.url}: {ref.base} is not configured; add it under "
                     "code_hosts to fetch this link")  # fmt: skip
        elif ref.sha:  # not in the source tree: take the diff from the host
            diff = HostClient(host).diff(ref)
            raw = _diff_patch(ctx, diff, f"{ref.repo}/{ref.sha}", ref.sha[:12], ref.url,
                              "jira_link", keys, hint)  # fmt: skip
            found += [raw] if raw else []
        else:
            client = HostClient(host)
            pr = client.pull_request(ref)
            local = [r for sha, _ in pr.commits if (r := _local_commit(ctx, sha, hint))]
            if local:
                found += local
            else:  # squashed or rebased on merge: the PR's diff stands for it
                text = f"{pr.title}\n\n{pr.body}\n\n{ref.url}".strip()
                raw = _diff_patch(ctx, client.diff(ref), f"{ref.repo}/pr-{ref.number}",
                                  pr.title or ref.url, text, "jira_link", keys, hint)  # fmt: skip
                found += [raw] if raw else []
    except (CodeHostError, KeyError, OSError) as e:
        ctx.problem(keys, f"jira_query: {ref.url}: {e}")
    for raw in found:
        raw.jira_keys = list(dict.fromkeys([*raw.jira_keys, *keys]))
    return found


ADAPTERS = {
    "git_range": from_git_range,
    "diff_folder": from_diff_folder,
    "spreadsheet": from_spreadsheet,
    "jira_query": from_jira_query,
}
